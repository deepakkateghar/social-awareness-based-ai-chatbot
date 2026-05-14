import streamlit as st
from openai import OpenAI
from docx import Document
from io import BytesIO
import time
from dotenv import load_dotenv
import os

load_dotenv()

st.set_page_config(
    page_title="AI Content Generator",
    page_icon="🤖",
    layout="wide",
    initial_sidebar_state="expanded"
)

st.markdown("""
<style>
html, body, [class*="css"] {
    font-family: 'Segoe UI', sans-serif;
}
.main {
    padding-top: 1rem;
}
.block-container {
    padding-top: 2rem;
}
.stTextInput > div > div > input {
    border-radius: 14px;
    padding: 14px;
    font-size: 16px;
}
.stButton > button {
    width: 100%;
    height: 3.2rem;
    border-radius: 14px;
    font-size: 18px;
    font-weight: 600;
}
.result-container {
    background-color: #111827;
    padding: 30px;
    border-radius: 18px;
    border: 1px solid #374151;
    margin-top: 20px;
}
.header-box {
    padding: 20px;
    border-radius: 18px;
    background: linear-gradient(to right, #111827, #1F2937);
    border: 1px solid #374151;
    margin-bottom: 25px;
}
</style>
""", unsafe_allow_html=True)

st.markdown("""
<div class="header-box">
    <h1>Social Awareness AI Chatbot</h1>
</div>
""", unsafe_allow_html=True)

with st.sidebar:

    st.header("Generation Custom Settings")

    tone = st.selectbox(
        "Writing Tone",
        [
            "Professional",
            "Academic",
            "Technical",
            "Conversational",
            "Creative"
        ]
    )
    detail_level = st.select_slider(
        "Detail Level",
        options=[
            "Basic",
            "Detailed",
            "Very Detailed",
            "Research Level"
        ],
        value="Very Detailed"
    )
    include_examples = st.checkbox(
        "Include Real-world Examples",
        value=True
    )

    include_case_studies = st.checkbox(
        "Include Case Studies",
        value=True
    )

    st.divider()

API_KEY = os.getenv("API_KEY")

client = OpenAI(
    api_key=API_KEY,
    base_url="https://openrouter.ai/api/v1"
)

col1, col2 = st.columns([4, 1])

with col1:

    topic = st.text_input(
        "Enter a topic to Query With",
        placeholder="Ask Anything..."
    )

with col2:
    st.write("")
    st.write("")
    generate = st.button("Ask")

if generate:
    if topic.strip() == "":
        st.warning("Please enter a topic.")
        st.stop()
    progress_bar = st.progress(0)
    loading_text = st.empty()
    loading_steps = [
        "Analyzing topic...",
        "Preparing structure...",
        "Generating introduction...",
        "Writing detailed sections...",
        "Adding examples and insights...",
        "Finalizing article..."
    ]

    for i, step in enumerate(loading_steps):
        loading_text.info(step)
        progress_bar.progress((i + 1) * 15)
        time.sleep(0.5)

    prompt = f"""
You are an expert professional content writer and researcher.

Write a HIGHLY DETAILED, ORIGINAL, HUMAN-LIKE article.

TOPIC:
{topic}

WRITING STYLE:
- Tone: {tone}
- Detail Level: {detail_level}

IMPORTANT REQUIREMENTS:
- Generate 3500+ words
- Human-like writing
- Natural flow
- Deep explanations
- Practical examples
- Professional readability
- No robotic tone
- No repetitive phrasing
- Well-structured sections
- Realistic insights
- Academic-quality content

INCLUDE:
- Introduction
- History/Evolution
- Core Concepts
- Applications
- Advantages
- Disadvantages
- Challenges
- Future Scope
- Conclusion

EXTRA REQUIREMENTS:
- Include examples: {include_examples}
- Include case studies: {include_case_studies}

FORMATTING:
- Proper headings
- Proper paragraphs
- Clean readable structure
- No markdown symbols
"""

    try:

        response = client.chat.completions.create(

            model="meta-llama/llama-3-8b-instruct",

            messages=[
                {
                    "role": "system",
                    "content": "You are a professional article writer."
                },
                {
                    "role": "user",
                    "content": prompt
                }
            ],

            temperature=0.7,

            max_tokens=4000
        )

        generated_text = response.choices[0].message.content

        progress_bar.progress(100)

        loading_text.success("Response Generated")
        tab1, tab2 = st.tabs(
            [
                "Preview",
                "Export Document"
            ]
        )

        with tab1:
            st.markdown("## Generated Response")
            st.markdown(
                f"""
                <div class="result-container">
                {generated_text}
                </div>
                """,
                unsafe_allow_html=True
            )

        with tab2:

            st.markdown("## Export Options")


            document = Document()

            document.add_heading(topic, level=1)

            for para in generated_text.split("\n"):

                if para.strip():

                    document.add_paragraph(para)

            buffer = BytesIO()

            document.save(buffer)

            buffer.seek(0)

            st.download_button(
                label="Download DOCX",
                data=buffer,
                file_name=f"{topic}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )

            # TXT

            st.download_button(
                label="Download TXT",
                data=generated_text,
                file_name=f"{topic}.txt",
                mime="text/plain"
            )

        st.markdown("---")

        st.caption(
            "Built with Streamlit by Kateghar Deepak."
        )

    except Exception as e:

        st.error(f"Error: {e}")